# -*- coding: utf-8 -*-
"""Экспорт текущего списка контрактов в .xlsx (для отправки/печати/архива)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from calculations import calc_profit, calc_margin_pct
import db

HEADERS = [
    "№", "Площадка (ЭТП)", "Заказчик", "Номер контракта", "Реестровая запись", "Дата заключения контракта",
    "По какому закону", "Товары в контракте", "Кол-во товара (всего)",
    "Сумма контракта, руб.", "Себестоимость, руб.", "Стоимость логистики, руб.",
    "Комиссия площадки, руб.", "Другие расходы, руб.", "Сумма обеспечения/гарантии, руб.",
    "Налог по контракту (справочно), руб.", "Прибыль по контракту (справочно), руб.", "% рентабельности",
    "Статус контракта", "Подписать до", "Срок исполнения", "Дата вручения продукции заказчику",
    "Статус оплаты", "Крайний срок оплаты", "Статус исполнения",
    "Ответственный за закупку", "Ответственный за получение", "Примечание",
]


def export_to_excel(purchases, filepath: str, conn=None):
    """Экспорт контрактов и точных месячных итогов при переданном соединении БД."""
    if conn is None:
        raise ValueError("Для корректного экспорта налогов требуется соединение с базой данных")
    wb = Workbook()
    ws = wb.active
    ws.title = "Контракты"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E75B6")
    border = Border(*(Side(style="thin", color="B7B7B7"),) * 4)
    money_fmt = '#,##0.00" ₽"'
    pct_fmt = "0.0%"
    date_fmt = "DD.MM.YYYY"

    for col, text in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=text)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.freeze_panes = "A2"

    money_cols = {10, 11, 12, 13, 14, 15, 16, 17}
    date_cols = {6, 20, 21, 22, 24}

    row_idx = 2
    for i, p in enumerate(purchases, start=1):
        contract_sum = p["contract_sum"] or 0.0
        # Налог на уровне отдельного контракта — справочный: общие расходы
        # месяца и общая налоговая база учитываются в листе «Итоги по месяцам».
        period = db._summary_period_for_purchase(conn, p) if conn is not None else None
        if conn is not None and period is not None:
            group = {k: float(p.get(k) or 0) for k in (
                "contract_sum", "purchase_cost", "logistics", "commission", "other_costs", "guarantee")}
            tax, _, _ = db._calc_month_tax(conn, period.year, period.month, group)
            profit = calc_profit(contract_sum, p["purchase_cost"] or 0.0, p["logistics"] or 0.0,
                                  p["commission"] or 0.0, p["other_costs"] or 0.0,
                                  p["guarantee"] or 0.0, tax)
            margin = calc_margin_pct(profit, contract_sum)
        else:
            tax = profit = margin = None
        resp_purchase = ", ".join(filter(None, [p.get("resp_purchase_name"), p.get("resp_purchase_phone"),
                                                 p.get("resp_purchase_email")]))
        resp_receiving = ", ".join(filter(None, [p.get("resp_receiving_name"), p.get("resp_receiving_phone"),
                                                  p.get("resp_receiving_email")]))

        values = [
            i, p["platform"], p["customer"], p["contract_no"], p.get("registry_record"), p["contract_date"],
            p["law"], p.get("product", ""), p.get("qty", 0), contract_sum, p["purchase_cost"],
            p["logistics"], p["commission"], p["other_costs"], p["guarantee"],
            tax, profit, margin, p["contract_status"], p.get("sign_deadline"), p["deadline"],
            p["handover_date"], p["payment_status"], p.get("payment_deadline"), p["exec_status"],
            resp_purchase, resp_receiving, p["note"],
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border = border
            if col in money_cols:
                cell.number_format = money_fmt
            elif col == 18:
                cell.number_format = pct_fmt
            elif col in date_cols:
                cell.number_format = date_fmt
        row_idx += 1

    widths = [5, 20, 20, 16, 20, 16, 14, 30, 12, 14, 14, 14, 13, 13, 15, 13, 14, 12, 16, 14, 13, 16, 14, 14,
              14, 24, 24, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{row_idx - 1}"
    if conn is not None:
        summary = wb.create_sheet("Итоги по месяцам")
        summary.append(["Год", "Месяц", "Контрактов", "Сумма контрактов", "Себестоимость",
                        "Логистика", "Комиссии", "Расходы контрактов", "Обеспечение",
                        "Прочие расходы месяца", "Налоговый режим", "Ставка, %",
                        "Налоговая база", "Налог", "Чистая прибыль", "Рентабельность"])
        for item in db.monthly_summary(conn):
            summary.append([item["year"], item["month"], item["contracts_count"],
                            item["contract_sum"], item["purchase_cost"], item["logistics"],
                            item["commission"], item["other_costs"], item["guarantee"],
                            item["monthly_expenses"], item["tax_regime"], item["tax_rate"],
                            item["tax_base"], item["tax"], item["profit"], item["margin_pct"]])
        for row in summary.iter_rows(min_row=2):
            for cell in row:
                if cell.column in (4, 5, 6, 7, 8, 9, 10, 13, 14, 15):
                    cell.number_format = money_fmt
                if cell.column == 16:
                    cell.number_format = pct_fmt
        for cell in summary[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True)
        summary.freeze_panes = "A2"
        summary.auto_filter.ref = summary.dimensions
        for idx in range(1, 17):
            summary.column_dimensions[get_column_letter(idx)].width = 22
        summary.column_dimensions["K"].width = 26
        summary.append([])
        summary.append(["Итоги месяца рассчитаны штатной функцией приложения; справочные значения по отдельным контрактам не включают распределение общих расходов месяца."])
    wb.save(filepath)