"""Регрессия экспорта: ставка месяца и точные месячные итоги без двойного учёта."""
import os
import sys
import tempfile
from openpyxl import load_workbook
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'app'))
import db
from excel_export import export_to_excel


def test_export_matches_monthly_summary_and_tax_regime():
    with tempfile.TemporaryDirectory() as tmp:
        conn = db.get_connection(os.path.join(tmp, 'test.db'))
        db.set_tax_profile(conn, '2026-10-01', 'Доходы минус расходы', 15)
        header = dict(customer='QA', contract_no='EXCEL-1', created_at='2026-10-08 11:00:00',
                      contract_date='2026-10-08', contract_sum=100000, purchase_cost=60000,
                      contract_status='Заключен', exec_status='В процессе',
                      payment_status='Не оплачено')
        db.insert_purchase(conn, header, [])
        path = os.path.join(tmp, 'export.xlsx')
        export_to_excel(db.fetch_all(conn), path, conn=conn)
        wb = load_workbook(path, data_only=True)
        assert wb['Контракты']['P2'].value == 6000
        assert wb['Контракты']['Q2'].value == 34000
        monthly = next(x for x in db.monthly_summary(conn) if x['year'] == 2026 and x['month'] == 10)
        sheet = wb['Итоги по месяцам']
        assert sheet['N2'].value == monthly['tax'] == 6000
        assert sheet['O2'].value == monthly['profit'] == 34000
        assert sheet['K2'].value == 'Доходы минус расходы'
        conn.close()


def test_monthly_shared_expenses_not_allocated_to_single_contract():
    with tempfile.TemporaryDirectory() as tmp:
        conn = db.get_connection(os.path.join(tmp, 'test.db'))
        db.set_tax_profile(conn, '2026-10-01', 'Доходы минус расходы', 15)
        db.insert_purchase(conn, dict(customer='QA', contract_no='E2', contract_date='2026-10-08',
                        created_at='2026-10-08 10:00:00', contract_sum=100000,
                        purchase_cost=60000), [])
        # Прочие расходы месяца влияют только на итоговую агрегацию.
        conn.execute('INSERT INTO monthly_expenses (expense_date, amount, category, description, period_year, period_month) VALUES (?, ?, ?, ?, ?, ?)',
                     ('2026-10-08', 10000, 'Прочее', 'Тест', 2026, 10))
        conn.commit()
        path = os.path.join(tmp, 'export.xlsx')
        export_to_excel(db.fetch_all(conn), path, conn=conn)
        wb = load_workbook(path, data_only=True)
        assert wb['Контракты']['P2'].value == 6000  # справочный расчёт
        assert wb['Итоги по месяцам']['N2'].value == 4500  # налог после общего расхода
        assert wb['Итоги по месяцам']['O2'].value == 25500
        conn.close()


if __name__ == '__main__':
    test_export_matches_monthly_summary_and_tax_regime()
    test_monthly_shared_expenses_not_allocated_to_single_contract()
    print('PASS: 2 Excel tax export tests')
